from pathlib import Path

import pytest
from docx import Document

from app.infrastructure.ingest.extract import SUPPORTED_EXTENSIONS, extract_text


def test_extracts_utf8_and_gb18030_text(tmp_path: Path) -> None:
    utf8 = tmp_path / "notes.md"
    utf8.write_text("视觉 Transformer\nViT", encoding="utf-8")
    legacy = tmp_path / "notes.txt"
    legacy.write_bytes("中文检索测试".encode("gb18030"))

    assert "Transformer" in extract_text(utf8)
    assert extract_text(legacy) == "中文检索测试"


def test_images_are_not_accepted(tmp_path: Path) -> None:
    image = tmp_path / "scan.png"
    image.write_bytes(b"not-an-image")

    assert ".png" not in SUPPORTED_EXTENSIONS
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(image)


def test_extracts_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "research.docx"
    document = Document()
    document.add_paragraph("Transformer 研究笔记")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "模型"
    table.cell(0, 1).text = "ViT"
    document.save(path)

    extracted = extract_text(path)

    assert "Transformer 研究笔记" in extracted
    assert "模型" in extracted
    assert "ViT" in extracted


def test_rejects_broken_docx(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")

    with pytest.raises(ValueError, match="DOCX"):
        extract_text(path)
